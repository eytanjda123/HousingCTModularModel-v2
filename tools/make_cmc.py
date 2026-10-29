# Builds the "City Models & Charts" export template (sub/cmc_template.xlsx) and its slot map (sub/cmc_map.json).
# Every slot cell holds a placeholder; the explorer writes values into slots with JS expressions (see cmc_map.json).
import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.label import DataLabelList

A='Arial'
F_T=Font(name=A,bold=True,size=14,color='1E2A32'); F_S=Font(name=A,bold=True,size=11,color='FFFFFF')
F_B=Font(name=A,bold=True); F_N=Font(name=A); F_I=Font(name=A,italic=True,color='595959'); F_G=Font(name=A,color='006100')
FILL_S=PatternFill('solid',fgColor='1E2A32'); FILL_K=PatternFill('solid',fgColor='E1E5E8'); FILL_H=PatternFill('solid',fgColor='DFE3E6')
BOT=Border(bottom=Side(style='thin',color='C9D1D6'))
FMT={'usd':'$#,##0','pct':'0.0%','pct0':'0%','num':'#,##0','x':'0.00','dec':'0.0000','mill':'0.00','txt':'@','yr':'0'}
CITIES=["Hartford","New Haven","Waterbury"]
CCOL={"Hartford":"2C6A8A","New Haven":"B5452F","Waterbury":"6B7F3A"}
SLOTS=[]   # {s,r,e,c,t}
def slot(ws,ref,expr,city=None,fmt='usd',font=F_N,bold=False):
    c=ws[ref]; c.value='' if fmt=='txt' else 0; c.number_format=FMT.get(fmt,fmt); c.font=F_B if bold else font
    SLOTS.append({"s":ws.title,"r":ref,"e":expr,"c":city,"t":"s" if fmt=='txt' else "n"})
def q(n): return "'"+n+"'" if ' ' in n else n

wb=openpyxl.Workbook()
# ---------------- Read Me
ws=wb.active; ws.title='Read Me'; ws.column_dimensions['A'].width=2; ws.column_dimensions['B'].width=130
ws['B2']='HOUSING CT, PREFAB MODULAR TEAM: city models and deck charts'; ws['B2'].font=F_T
slot(ws,'B3','STAMP',fmt='txt',font=F_I)
slot(ws,'B4','LINK',fmt='txt',font=F_I)
readme=["",
("TABS",F_B),
"Inputs Snapshot: every explorer input as it was set at export, for each city, with its status (sourced, placeholder, team choice) and note.",
"Hartford / New Haven / Waterbury: one building per city. Sections: 1 scenario, 2 city data, 3 key assumptions, 4 calculations (cost, rent, buyer capacity),",
"   5 results incl. both gaps, the funding stack and the owner's net monthly cost, 6 uses vs. sources, 7 funding waterfall, 8 owner-income sensitivity, 9 building size, 10 what moves the remaining gap.",
"Compare: the three cities side by side (linked to the city tabs).",
"Market Data: renter households by income band and the city profiles; rows the explorer controls (AMI, rents, sales, rent limit) come from the export.",
"Charts: the deck charts as native Excel charts, fed by data blocks linked to the other tabs.",
"",
("WHERE THE NUMBERS COME FROM",F_B),
"Every result is calculated by Prefab Gap Explorer 2 with the inputs set the way they were at export, including any city-only overrides. The Compare and Charts tabs link to the city tabs.",
"This workbook is a presentation snapshot: the city tabs hold values, not formulas. For the fully live model with every formula, use 'Download Excel model' in the explorer's Export dialog.",
"",
("HOW THE MODEL WORKS",F_B),
"1. Cost: modular factory + field cost per gross sq ft (falling with order volume), transport, site work, contingency, soft costs, construction loan, carry, land and the developer fee give the total development cost (TDC).",
"2. Rent: the rental unit charges the lower of the AMI rent (rent + utilities at the set share of a renter household's income, less the utility allowance) and market asking rent.",
"3. Buyer: the lower of the lender test (40% of own income + 75% of rent) and the owner test (30% of own income + rent collected), after tax, insurance, utilities and reserves, plus CHFA Time To Own.",
"4. Appraisal: the base case is a size-adjusted new-build value; the sales gap (appraisal − what the buyer can pay) is a DOH homebuyer subsidy, and the developer gap (TDC − sale proceeds) is filled by HDF, AHP, HTCC, NMTC (if on) and DOH.",
"5. Program amounts are caps and precedents, not awards. Whether the programs can fund the same home is for the law team to confirm.",
"",
("CHANGES FROM THE EXPLORER DEFAULTS",F_B)]
r=6
for x in readme:
    t,f=(x if isinstance(x,tuple) else (x,F_N)); ws.cell(r,2,t).font=f; r+=1
README_LOG=r   # change log rows start here (written by JS)
ws.cell(r,2,'').font=F_N

# ---------------- Inputs Snapshot (rows appended by JS)
ws=wb.create_sheet('Inputs Snapshot')
ws['B2']='INPUTS SNAPSHOT: every explorer input at export'; ws['B2'].font=F_T
ws['B3']='Values shown as the explorer formats them. "City only" marks a value set for one city.'; ws['B3'].font=F_I
for j,h in enumerate(["Input","Group","Hartford","New Haven","Waterbury","Status","Note"],2):
    c=ws.cell(5,j,h); c.font=F_S; c.fill=FILL_S
for j,w in enumerate([2,52,24,18,18,18,16,110],1): ws.column_dimensions[L(j)].width=w
ws['B6']=''; ws['B6'].font=F_N  # style anchor
SNAP_ROW=6

# ---------------- City tabs
ROWS={}   # key -> row (same on each city tab)
def city_tab(city):
    ws=wb.create_sheet(city)
    for j,w in enumerate([2,58,16,16,16,16,16,16,16,40],1): ws.column_dimensions[L(j)].width=w
    ws['B2']=f'{city.upper()}: one building, prefab modular'; ws['B2'].font=F_T
    slot(ws,'B3','"Exported from Prefab Gap Explorer 2. Design: "+designOf(T)+". Values, not formulas: the explorer computed every number with the inputs set at export."',city,'txt',F_I)
    r=[5]
    def sec(t):
        c=ws.cell(r[0],2,t); c.font=F_S; c.fill=FILL_S
        for j in range(3,10): ws.cell(r[0],j).fill=FILL_S
        r[0]+=1
    def row(label,key,expr,fmt='usd',note='',bold=False):
        ws.cell(r[0],2,label).font=F_B if bold else F_N
        slot(ws,f'C{r[0]}',expr,city,fmt,bold=bold)
        if note: ws.cell(r[0],5 if fmt!='txt' else 6,note).font=F_I
        if bold:
            for j in (2,3): ws.cell(r[0],j).fill=FILL_K
        ROWS[key]=r[0]; r[0]+=1
    def gap(): r[0]+=1
    sec('1. SCENARIO')
    row('City','city','city','txt')
    row('Building design','design','designOf(T)','txt')
    row('Total units in the building','units','s.units','num')
    row('Owner unit bedrooms','ownerBR','s.ownerBR','txt')
    row('Rental unit bedrooms','rentBR','s.rentBR','txt')
    row('Owner household income target','ownerAMI','s.ownerAMI','pct0',"% of AMI")
    row('Owner household size','hhSize','s.hhSize','num')
    row('Rental unit rent target','rentAMI','s.rentAMI','pct0','% of AMI; 0% = market rent')
    row('Cap the AMI rent at market rent?','capMkt','s.capMkt?"Yes":"No"','txt')
    row('Construction method','method','s.method','txt')
    row('Order volume (buildings ordered together)','volume','s.volume','num')
    row('Appraisal basis','basis','(BASIS.find(b=>b[0]===s.basis)||[0,s.basis])[1]','txt')
    row('Property tax assessed on','taxBasis','(TAXB.find(b=>b[0]===s.taxBasis)||[0,s.taxBasis])[1]','txt')
    gap(); sec('2. CITY DATA')
    row('HUD 50% income limit, 4-person','ami50','c.ami50','usd','HUD FY2026 Section 8 limits')
    row('100% AMI used in the model (2 x 50% limit)','ami100','R.ami100')
    ws.cell(r[0],2,'Bedrooms').font=F_I
    for j,b in enumerate(["Studio","1BR","2BR","3BR","4BR"]): c=ws.cell(r[0],3+j,b); c.font=F_B; c.fill=FILL_H
    r[0]+=1
    ws.cell(r[0],2,'Market asking rent by bedroom count ($ per month)').font=F_N
    for j in range(5): slot(ws,f'{L(3+j)}{r[0]}',f'c.mkt[{j}]',city)
    ROWS['mktRow']=r[0]; r[0]+=1
    ws.cell(r[0],2,'Building size').font=F_I
    for j,b in enumerate(["2-unit","3-unit","4-6 unit"]): c=ws.cell(r[0],3+j,b); c.font=F_B; c.fill=FILL_H
    r[0]+=1
    ws.cell(r[0],2,'Median sale, existing buildings (CT CAMA 2023-26)').font=F_N
    for j in range(3): slot(ws,f'{L(3+j)}{r[0]}',f'c.med[{j}]',city)
    ROWS['medRow']=r[0]; r[0]+=1
    row('Size-adjusted new-build value (base case)','custom','c.custom')
    row('DOH new-build precedent','prec','c.prec')
    row('Assessment ratio used','assess','R.assessUsed','pct')
    row('Mill rate','mill','c.mill','mill','$ per $1,000 assessed')
    row('Effective property tax rate','effTax','R.assessUsed*c.mill/1000','pct')
    row('Homeowner insurance ($ per unit per year)','ins','c.ins')
    row('Land acquisition and closing ($ per building)','land','c.land')
    gap(); sec('3. KEY ASSUMPTIONS (full list on Inputs Snapshot)')
    for lab,k,e,f in [('Mortgage rate','rate','s.rate','pct'),('Mortgage term (years)','term','s.term','num'),('Down payment','dp','s.dp','pct'),('FHA mortgage insurance (annual)','mip','s.mip','pct'),
        ('Lender credit on rental income','rentCredit','s.rentCredit','pct0'),('Lender test cap, % of qualifying income','lenderCap','s.lenderCap','pct0'),('Owner test cap, % of own income','ownerCap','s.ownerCap','pct0'),
        ('Rent + utilities as % of renter income','rentShare','s.rentShare','pct'),('Tenant utility allowance (rental unit)','uaR','s.ua[BR.indexOf(s.rentBR)]','usd'),('Vacancy and collection loss','vac','s.vac','pct'),
        ('Maintenance and reserve (% of value per year)','maintPct','s.maintPct','pct'),('Owner energy bills ($ per month)','energy','s.energy','usd'),('Owner-paid common utilities ($ per unit per month)','common','s.common','usd'),
        ('Gross-up (net to gross sq ft)','grossup','s.grossup','x'),('Hard-cost contingency','hardCont','s.hardCont','pct'),('Soft-cost contingency','softCont','s.softCont','pct'),
        ('Construction loan rate','loanRate','s.loanRate','pct'),('Developer fee (% of cost)','fee','s.fee','pct')]:
        row(lab,k,e,f)
    gap(); sec('4. CALCULATIONS')
    ws.cell(r[0],2,'Building and development cost').font=F_B; r[0]+=1
    for lab,k,e,f,b in [('Owner unit net sq ft','ownerSF','R.ownerSF','num',0),('Rental unit net sq ft (each)','rentSF','R.rentSF','num',0),('Rented units','rented','R.rented','num',0),
        ('Gross sq ft, whole building','gsf','R.gsf','num',0),('Factory cost per gross sq ft (at this order volume)','facPSF','R.facPSF','usd',0),('Field cost per gross sq ft','fieldPSF','R.fieldPSF','usd',0),
        ('Transport, crane and set','trans','R.trans','usd',0),('Building (factory + field + transport, or stick-built)','bldg','R.bldg','usd',0),('Add-ons and hard-cost extras','feat','R.featCost+R.hx','usd',0),
        ('Warehouse and bulk buying (net)','wh','R.whNet','usd',0),('Site work','site','R.site','usd',0),('Hard-cost contingency','hardContC','R.hardContC','usd',0),
        ('Construction cost','construction','R.construction','usd',1),('Soft costs incl. contingency','soft','R.soft','usd',0),('Construction loan interest','interest','R.interest','usd',0),
        ('Carry and land','carryLand','R.carry+R.land','usd',0),('Developer fee','feeC','R.feeC','usd',0),('Total development cost (TDC)','tdc','R.tdc','usd',1),
        ('TDC per unit','tdcUnit','R.tdcUnit','usd',0),('TDC per gross sq ft','tdcPSF','R.tdcPSF','usd',0)]:
        row(lab,k,e,f,bold=bool(b))
    gap(); ws.cell(r[0],2,'Rent for the rental unit').font=F_B; r[0]+=1
    for lab,k,e,f in [('Renter household income at the rent target','tenantInc','R.tenantInc','usd'),('AMI rent, net of utilities (what the owner can charge)','rentNet','R.rentNet','usd'),
        ('Market asking rent, rental bedroom count','market','R.market','usd'),('Rent charged per rental unit','rent','R.rent','usd'),('What sets the rent','rentBinding','R.rentBinding','txt'),
        ('Rent charged as % of renter income','rentPctInc','R.rentPctInc','pct'),('Gross rent, all rental units, per month','grossRent','R.grossRent','usd'),
        ('Rent collected per month (after vacancy)','collected','R.collected','usd'),('Rent a lender counts, per month','lenderRentM','R.lenderRentM','usd')]:
        row(lab,k,e,f)
    gap(); ws.cell(r[0],2,'What the buyer can pay').font=F_B; r[0]+=1
    for lab,k,e,f in [('Owner household income','income','R.income','usd'),('Property tax per month','tax','R.tax','usd'),('Insurance per month','insM','R.insM','usd'),
        ('Owner-paid common utilities per month','commonM','R.commonM','usd'),('Owner energy bills per month','energyM','R.energyM','usd'),('Maintenance and reserve per month','maintM','R.maintM','usd'),
        ('Less add-on savings per month','featSave','-R.featSaveMo','usd'),('Non-mortgage housing costs per month','opex','R.opex','usd'),
        ('Monthly payment per $1 of loan (P&I + FHA insurance)','k','R.k','dec'),('Lender test: maximum price','testA','R.testA','usd'),('Owner test: maximum price','testB','R.testB','usd'),
        ('Buyer capacity (lower of the two tests)','capacity','R.capacity','usd'),('Which test binds','capBinding','R.capBinding','txt'),('CHFA Time To Own and other buyer assistance','tto','R.tto','usd'),
        ('Most the buyer can pay','afford','R.afford','usd')]:
        row(lab,k,e,f,bold=(k=='afford'))
    gap(); sec('5. RESULTS')
    for lab,k,e,f,b in [('Appraised value','fmv','R.fmv','usd',0),('Sales gap (appraisal − what the buyer can pay)','salesGap','R.salesGap','usd',1),('DOH homebuyer subsidy','hb','R.hb','usd',0),
        ('Sale proceeds to the developer','proceeds','R.proceeds','usd',0),('What the buyer pays (mortgage + down payment)','buyerPays','R.buyerPays','usd',0),
        ('Developer gap (TDC − sale proceeds)','devGap','R.devGap','usd',1),('HDF / C4C grant','sHdf','R.sHdf','usd',0),('FHLBank Boston AHP','sAhp','R.sAhp','usd',0),('HTCC share','sHtcc','R.sHtcc','usd',0),
        ('New Markets Tax Credits (net, if on)','sNmtc','R.sNmtc','usd',0),('DOH developer subsidy','sDoh','R.sDoh','usd',0),('Remaining developer gap','remain','R.remain','usd',1),
        ('Total subsidy needed (TDC − what the buyer pays)','subNeed','R.subNeed','usd',1),('Subsidy needed per unit','subUnit','R.subNeed/s.units','usd',0),('Subsidy as % of TDC','subPct','R.subNeed/R.tdc','pct',0),
        ('Most the buyer can pay as % of the existing-building median','priceVsMed','R.afford/c.med[R.bucket]','pct',0)]:
        row(lab,k,e,f,bold=bool(b))
    gap(); ws.cell(r[0],2,"Owner's net monthly housing cost").font=F_B; r[0]+=1
    for lab,k,e,f,b in [('Mortgage payment + FHA insurance','pi','R.pi','usd',0),('Plus non-mortgage costs (tax, insurance, utilities, reserve, less add-on savings)','opex2','R.opex','usd',0),
        ('Total monthly cost of owning the building','ownTotal','R.ownTotal','usd',0),('Minus rent collected (after vacancy)','collected2','R.collected','usd',0),
        ("Owner's net monthly cost",'net','R.net','usd',1),('Net cost as % of owner income','netPct','R.netPct','pct',1),('Residual income after housing and MIT living costs ($ per year)','residual','R.residual','usd',0),
        ('Energy burden','energyBurden','R.energyBurden','pct',0)]:
        row(lab,k,e,f,bold=bool(b))
    gap(); sec('6. USES VS. SOURCES (chart data)')
    for j,h in enumerate(['Uses','Sources'],3): c=ws.cell(r[0],j,h); c.font=F_B; c.fill=FILL_H
    r[0]+=1; ROWS['usStart']=r[0]
    US=[('Building','R.bldg','0'),('Add-ons and warehouse','R.featCost+R.hx+R.whNet','0'),('Site work and hard contingency','R.site+R.hardContC','0'),('Soft costs','R.soft','0'),
        ('Construction loan, carry and land','R.interest+R.carry+R.land','0'),('Developer fee','R.feeC','0'),('Buyer down payment','0','R.down'),('Buyer mortgage','0','R.mortgage'),
        ('Buyer assistance (Time To Own)','0','R.tto'),('DOH homebuyer subsidy','0','R.hb'),('HDF / C4C grant','0','R.sHdf'),('FHLBank Boston AHP','0','R.sAhp'),('HTCC share','0','R.sHtcc'),
        ('New Markets Tax Credits','0','R.sNmtc'),('DOH developer subsidy','0','R.sDoh'),('Remaining gap','0','R.remain')]
    for lab,u,sx in US:
        ws.cell(r[0],2,lab).font=F_N; slot(ws,f'C{r[0]}',u,city); slot(ws,f'D{r[0]}',sx,city); r[0]+=1
    ROWS['usEnd']=r[0]-1
    ws.cell(r[0],2,'Total').font=F_B; ws.cell(r[0],3,f"=SUM(C{ROWS['usStart']}:C{ROWS['usEnd']})").number_format=FMT['usd']; ws.cell(r[0],4,f"=SUM(D{ROWS['usStart']}:D{ROWS['usEnd']})").number_format=FMT['usd']
    for j in (3,4): ws.cell(r[0],j).font=F_B
    ROWS['usTot']=r[0]; r[0]+=1
    ws.cell(r[0],2,'Check: sources minus uses (0, or a surplus when the sale covers the cost)').font=F_I; c=ws.cell(r[0],3,f"=D{ROWS['usTot']}-C{ROWS['usTot']}"); c.number_format=FMT['usd']; r[0]+=2
    sec('7. FUNDING WATERFALL (caps and precedents, not awards)')
    for j,h in enumerate(['Amount','Chart: base','Chart: bar'],3): c=ws.cell(r[0],j,h); c.font=F_B; c.fill=FILL_H
    r[0]+=1; w0=r[0]; ROWS['wfStart']=w0
    WF=[('Total development cost','R.tdc'),('Buyer mortgage + down payment','R.mortgage+R.down'),('Time To Own','R.tto'),('DOH homebuyer subsidy','R.hb'),('HDF / C4C grant','R.sHdf'),
        ('FHLBank Boston AHP','R.sAhp'),('HTCC share','R.sHtcc'),('New Markets Tax Credits','R.sNmtc'),('DOH developer subsidy','R.sDoh'),('Remaining gap','R.remain')]
    for i,(lab,e) in enumerate(WF):
        rr=w0+i; ws.cell(rr,2,lab).font=F_N; slot(ws,f'C{rr}',e,city)
        if i==0: ws.cell(rr,4,0); ws.cell(rr,5,f'=C{rr}')
        elif i==len(WF)-1: ws.cell(rr,4,0); ws.cell(rr,5,f'=C{rr}')
        else: ws.cell(rr,4,f'=MAX(0,$C${w0}-SUM($C${w0+1}:C{rr}))'); ws.cell(rr,5,f'=MIN(C{rr},MAX(0,$C${w0}-SUM($C${w0+1}:C{rr-1})))' if i>1 else f'=MIN(C{rr},$C${w0})')
        for j in (4,5): ws.cell(rr,j).number_format=FMT['usd']; ws.cell(rr,j).font=F_N
    ROWS['wfEnd']=w0+len(WF)-1; r[0]=w0+len(WF)
    ws.cell(r[0],2,'Program amounts are caps and precedents, not awards. Whether HDF, AHP, HTCC and DOH can fund the same home is unconfirmed.').font=F_I; r[0]+=2
    sec("8. SENSITIVITY: owner income target (everything else as set)")
    LV=[0.5,0.6,0.7,0.8,0.9,1.0,1.2]
    ws.cell(r[0],2,'Owner income, % of AMI').font=F_B
    for j,v in enumerate(LV): c=ws.cell(r[0],3+j,v); c.number_format=FMT['pct0']; c.font=F_B; c.fill=FILL_H
    ROWS['amiHead']=r[0]; r[0]+=1
    for lab,k,e,f in [('Owner household income','sIncome','income','usd'),('Lender test: maximum price','sTestA','testA','usd'),('Owner test: maximum price','sTestB','testB','usd'),
        ('Most the buyer can pay (incl. Time To Own)','sAfford','afford','usd'),('Appraised value','sFmv','fmv','usd'),('Sales gap','sSales','salesGap','usd'),
        ('Remaining developer gap','sRemain','remain','usd'),('Subsidy needed per unit','sSubU','subU','usd'),("Owner's net monthly cost",'sNet','net','usd'),('Net cost as % of income','sNetPct','netPct','pct')]:
        ws.cell(r[0],2,lab).font=F_N
        for j,v in enumerate(LV): slot(ws,f'{L(3+j)}{r[0]}',f'AMI({v}).{e}',city,f)
        ROWS[k]=r[0]; r[0]+=1
    r[0]+=1
    sec('9. BUILDING SIZE: modular vs. stick-built (same city data and inputs)')
    NS=[2,3,4]
    ws.cell(r[0],2,'Units in the building').font=F_B
    for j,v in enumerate(NS): c=ws.cell(r[0],3+j,v); c.font=F_B; c.fill=FILL_H
    ROWS['sizeHead']=r[0]; r[0]+=1
    for lab,k,e,f in [('Gross sq ft','zGsf','U(n,0).gsf','num'),('Modular TDC','zModTdc','U(n,0).tdc','usd'),('Modular TDC per unit','zModU','U(n,0).tdcUnit','usd'),
        ('Stick-built TDC','zStkTdc','U(n,1).tdc','usd'),('Stick-built TDC per unit','zStkU','U(n,1).tdcUnit','usd'),('Modular saving per unit','zSave','1-U(n,0).tdcUnit/U(n,1).tdcUnit','pct'),
        ('Appraised value used','zFmv','U(n,0).fmv','usd'),('Modular subsidy needed per unit','zSubU','U(n,0).subNeed/n','usd'),('Modular remaining developer gap','zRemain','U(n,0).remain','usd')]:
        ws.cell(r[0],2,lab).font=F_N
        for j,n in enumerate(NS): slot(ws,f'{L(3+j)}{r[0]}',e.replace('(n','('+str(n)).replace('/n','/'+str(n)),city,f)
        ROWS[k]=r[0]; r[0]+=1
    ws.cell(r[0],2,'Modular = the construction method set in the explorer (Reframe unless changed). The appraisal stays at the value set for the base building.').font=F_I; r[0]+=2
    sec('10. WHAT MOVES THE REMAINING GAP: one change at a time')
    for j,h in enumerate(['Remaining gap','Change vs. as set','Subsidy needed'],3): c=ws.cell(r[0],j,h); c.font=F_B; c.fill=FILL_H
    r[0]+=1; ROWS['scStart']=r[0]
    ws.cell(r[0],2,'As set at export').font=F_B; slot(ws,f'C{r[0]}','R.remain',city,bold=True); ws.cell(r[0],4,0).number_format=FMT['usd']; slot(ws,f'E{r[0]}','R.subNeed',city); r[0]+=1
    for i in range(NSCEN):
        slot(ws,f'B{r[0]}',f'SCEN[{i}].l',city,'txt'); slot(ws,f'C{r[0]}',f'SC({i}).remain',city)
        c=ws.cell(r[0],4,f"=C{r[0]}-$C${ROWS['scStart']}"); c.number_format='+$#,##0;-$#,##0;$0'; c.font=F_N
        slot(ws,f'E{r[0]}',f'SC({i}).subNeed',city); r[0]+=1
    ws.cell(r[0],2,'Each row reruns the explorer model with one change from the scenario as set.').font=F_I
    ws.freeze_panes='C5'
    return ws
NSCEN=21
for cty in CITIES: city_tab(cty)

# ---------------- Compare
ws=wb.create_sheet('Compare')
ws['B2']='THREE-CITY COMPARISON (linked to the city tabs)'; ws['B2'].font=F_T
for j,cty in enumerate(CITIES,3): c=ws.cell(4,j,cty); c.font=F_S; c.fill=FILL_S
ws.cell(4,2,'').fill=FILL_S
for j,w in enumerate([2,56,18,18,18],1): ws.column_dimensions[L(j)].width=w
CMP=[('Building design','design','txt'),('Owner income at target AMI','income','usd'),('Most the buyer can pay','afford','usd'),('Appraised value','fmv','usd'),
     ('Most the buyer can pay as % of the existing-building median','priceVsMed','pct'),('Sales gap','salesGap','usd'),("Owner's net monthly cost (after rent received)",'net','usd'),
     ('Net cost as % of owner income','netPct','pct'),('Rent charged, rental unit','rent','usd'),('Market asking rent, rental unit','market','usd'),('Total development cost','tdc','usd'),
     ('TDC per unit','tdcUnit','usd'),('Developer gap','devGap','usd'),('Remaining developer gap','remain','usd'),('Total subsidy needed','subNeed','usd'),('Subsidy needed per unit','subUnit','usd')]
CMPROW={}
for i,(lab,k,f) in enumerate(CMP):
    rr=5+i; ws.cell(rr,2,lab).font=F_B if k in('remain','subUnit') else F_N
    for j,cty in enumerate(CITIES,3):
        c=ws.cell(rr,j,f"={q(cty)}!$C${ROWS[k]}"); c.number_format=FMT[f]; c.font=F_N
    CMPROW[k]=rr

# ---------------- Market Data
ws=wb.create_sheet('Market Data')
ws['B2']='MARKET DATA: renter households by income band, 2024'; ws['B2'].font=F_T
ws['B3']='Source: team AMI Distribution Analysis (ACS 2024 1-year B25118 renter households by income, HUD FY2024 3-person limits, incomes spread evenly within brackets). Counts are estimates with ACS margins of error.'; ws['B3'].font=F_I
for j,w in enumerate([2,52,16,16,16,16,16,14],1): ws.column_dimensions[L(j)].width=w
BANDS=['Under 30% AMI','30-50% AMI','50-80% AMI','80-100% AMI','Over 100% AMI']
for j,h in enumerate(['Renter households']+BANDS+['Total'],2): c=ws.cell(5,j,h); c.font=F_S; c.fill=FILL_S
RH={"Hartford":[16254,6026,6832,2875,4437],"New Haven":[12538,8612,6594,3329,8313],"Waterbury":[12085,5163,4078,755,1542]}
for i,cty in enumerate(CITIES):
    ws.cell(6+i,2,cty).font=F_N
    for j,v in enumerate(RH[cty]): ws.cell(6+i,3+j,v).number_format='#,##0'
    ws.cell(6+i,8,f'=SUM(C{6+i}:G{6+i})').number_format='#,##0'
for j,h in enumerate(['Share of renter households']+BANDS,2): c=ws.cell(10,j,h); c.font=F_S; c.fill=FILL_S
for i,cty in enumerate(CITIES):
    ws.cell(11+i,2,cty).font=F_N
    for j in range(5): c=ws.cell(11+i,3+j,f'={L(3+j)}{6+i}/$H{6+i}'); c.number_format='0%'
ws['B16']="CITY PROFILES (deck: 'Three cities, three markets' slide)"; ws['B16'].font=F_B
ws['B17']='Sources: Census QuickFacts V2025; ACS 2024 1-year (homeownership B25118, median income B19013, poverty); Zillow asking rents Sept 2026 (New Haven) and Redfin Aug 2026; CT CAMA 2023-26; team AMI analysis. Rows marked "explorer" come from the export.'; ws['B17'].font=F_I
for j,h in enumerate(['Metric']+CITIES+['From'],2): c=ws.cell(18,j,h); c.font=F_S; c.fill=FILL_S
PROF=[('Population, July 2025',[121981,138774,115771],'num',None),('Population change since 2020',[0.003,0.036,0.012],'pct',None),('Homeownership rate',[0.26,0.3,0.46],'pct',None),
 ('Median household income (city)',[50418,59705,47881],'usd',None),('Area median income, 4-person (HUD published)',[129200,123200,114000],'usd',None),
 ('60% AMI, 4-person (model: 2 x 50% limit x 0.6)',None,'usd','2*c.ami50*0.6'),('80% AMI, 4-person (model)',None,'usd','2*c.ami50*0.8'),
 ('Asking rent, 1BR',None,'usd','c.mkt[1]'),('Asking rent, 2BR',None,'usd','c.mkt[2]'),('Asking rent, 3BR',None,'usd','c.mkt[3]'),
 ('60% AMI rent limit, 2BR (gross, rent + utilities)',None,'usd','2*c.ami50*s.brFactor[2]*0.6*s.rentShare/12'),('Median sale price, all homes (Redfin)',[339775,399735,309795],'usd',None),
 ('Median sale, existing 2-unit (CAMA)',None,'usd','c.med[0]'),('Renters below 50% AMI',[0.61,0.54,0.73],'pct0',None),('Poverty rate',[0.275,0.251,0.259],'pct',None),
 ('Hispanic or Latino (any race)',[0.445,0.31,0.385],'pct',None),('Black',[0.358,0.295,0.252],'pct',None),('White, non-Hispanic',[0.151,0.3,0.293],'pct',None),
 ('Income needed to buy the median home (% of AMI)',[1.16,1.2,1.15],'pct0',None),('Market role',['Deep-need rental','Growing, high-revenue','Ownership city'],'txt',None)]
MD={}
for i,(lab,vals,f,e) in enumerate(PROF):
    rr=19+i; ws.cell(rr,2,lab).font=F_N
    for j,cty in enumerate(CITIES):
        if e: slot(ws,f'{L(3+j)}{rr}',e,cty,f)
        else: c=ws.cell(rr,3+j,vals[j]); c.number_format=FMT[f]; c.font=F_N
    ws.cell(rr,6,'explorer' if e else 'static').font=F_I
    MD[lab]=rr
rr=19+len(PROF); ws.cell(rr,2,'City median income as % of AMI').font=F_N
for j in range(3): c=ws.cell(rr,3+j,f'={L(3+j)}{MD["Median household income (city)"]}/{L(3+j)}{MD["Area median income, 4-person (HUD published)"]}'); c.number_format='0%'
MD['medPct']=rr

# ---------------- Charts
ws=wb.create_sheet('Charts')
for j,w in enumerate([2,44,14,14,14,14,14,14,14],1): ws.column_dimensions[L(j)].width=w
ws['B2']='CHARTS: native Excel charts for the deck. Data blocks link to the city tabs and Market Data.'; ws['B2'].font=F_T
ws['B3']="To keep a chart linked in PowerPoint: copy it here and use Paste Special > 'Use Destination Theme & Link Data'. Right-click any series to recolor it."; ws['B3'].font=F_I
FOCUS_SLOTS=[]
def sty(ch,title,w=17,h=8.5):
    ch.title=title; ch.width=w; ch.height=h; ch.legend.position='b'
    try: ch.y_axis.majorGridlines.spPr=GraphicalProperties(ln={'solidFill':'E1E5E8'})
    except Exception: pass
    ch.x_axis.delete=False; ch.y_axis.delete=False
def color(ser,hexc,line=False):
    if line:
        ser.graphicalProperties.line.solidFill=hexc; ser.graphicalProperties.line.width=28000; ser.smooth=False
    else:
        ser.graphicalProperties.solidFill=hexc; ser.graphicalProperties.line.solidFill=hexc
def hide(ser): ser.graphicalProperties.noFill=True; ser.graphicalProperties.line.noFill=True
def blk(r,t): c=ws.cell(r,2,t); c.font=F_B
# A. affordability curve
r=5; blk(r,'A. Most a buyer can pay vs. the appraisal, by owner income (deck: affordability curve)')
r+=1; ws.cell(r,2,'Owner income, % of AMI').font=F_B
for j in range(7): c=ws.cell(r,3+j,f"={q('New Haven')}!{L(3+j)}{ROWS['amiHead']}"); c.number_format='0%'
hdrA=r
for i,cty in enumerate(CITIES):
    ws.cell(r+1+2*i,2,f'{cty}: most the buyer can pay').font=F_N; ws.cell(r+2+2*i,2,f'{cty}: appraised value').font=F_N
    for j in range(7):
        ws.cell(r+1+2*i,3+j,f"={q(cty)}!{L(3+j)}{ROWS['sAfford']}").number_format=FMT['usd']
        ws.cell(r+2+2*i,3+j,f"={q(cty)}!{L(3+j)}{ROWS['sFmv']}").number_format=FMT['usd']
ch=LineChart(); sty(ch,'Most a buyer can pay vs. the appraised value',20,9)
for i,cty in enumerate(CITIES):
    for k in (0,1):
        rr=hdrA+1+2*i+k; ch.add_data(Reference(ws,min_col=2,max_col=9,min_row=rr),from_rows=True,titles_from_data=True)
        color(ch.series[-1],CCOL[cty],True)
        if k: ch.series[-1].graphicalProperties.line.dashStyle='dash'
ch.set_categories(Reference(ws,min_col=3,max_col=9,min_row=hdrA)); ch.x_axis.title='Owner income, % of AMI'; ch.y_axis.numFmt='$#,##0'
ws.add_chart(ch,'B14')
# B. renter households
r=33; blk(r,'B. Renter households by income band, 2024 (deck: buyer pool)')
r+=1; ws.cell(r,2,'City').font=F_B
for j,b in enumerate(BANDS): ws.cell(r,3+j,b).font=F_B
for i,cty in enumerate(CITIES):
    ws.cell(r+1+i,2,f"='Market Data'!B{11+i}")
    for j in range(5): ws.cell(r+1+i,3+j,f"='Market Data'!{L(3+j)}{11+i}").number_format='0%'
ch=BarChart(); ch.type='bar'; ch.grouping='percentStacked'; ch.overlap=100; sty(ch,'Share of renter households by income band, 2024',20,7.5)
ch.add_data(Reference(ws,min_col=3,max_col=7,min_row=r,max_row=r+3),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=r+1,max_row=r+3))
for s_,cx in zip(ch.series,['3B4B55','7A8A92','B5452F','D9A07E','DFE3E6']): color(s_,cx)
ws.add_chart(ch,'F38')
# C. uses vs sources x3
STK=['1E2A32','3B4B55','7A8A92','A9B4BA','C9D1D6','9CC3D5','2C6A8A','5B8FA8','8FB8CC','D9A07E','6B7F3A','8E9F5E','B3C08A','D4AF37','B5452F','E06C55']
r=57
UCH=[]
for cty in CITIES:
    blk(r,f'C. Uses vs. sources, one {cty} building (deck: city model)')
    ws.cell(r+1,3,'Uses').font=F_B; ws.cell(r+1,4,'Sources').font=F_B
    n=ROWS['usEnd']-ROWS['usStart']+1
    for i in range(n):
        src=ROWS['usStart']+i
        ws.cell(r+2+i,2,f"={q(cty)}!B{src}"); ws.cell(r+2+i,3,f"={q(cty)}!C{src}").number_format=FMT['usd']; ws.cell(r+2+i,4,f"={q(cty)}!D{src}").number_format=FMT['usd']
    ch=BarChart(); ch.type='col'; ch.grouping='stacked'; ch.overlap=100; sty(ch,f'{cty}: uses vs. sources, one building',14,10)
    for i in range(n):
        ch.add_data(Reference(ws,min_col=2,max_col=4,min_row=r+2+i),from_rows=True,titles_from_data=True); color(ch.series[-1],STK[i%len(STK)])
    ch.set_categories(Reference(ws,min_col=3,max_col=4,min_row=r+1)); ch.y_axis.numFmt='$#,##0'; ch.legend.position='r'
    ws.add_chart(ch,f'F{r}')
    r+=22
# D. TDC per unit by building size (focus = New Haven tab)
blk(r,'D. TDC per unit, modular vs. stick-built (New Haven tab; deck: modular vs. stick)')
ws.cell(r+1,2,'Building size').font=F_B
for j in range(3): ws.cell(r+1,3+j,f"={q('New Haven')}!{L(3+j)}{ROWS['sizeHead']}&\" units\"")
for i,(lab,k) in enumerate([('Modular','zModU'),('Stick-built','zStkU'),('Modular saving','zSave'),('Modular subsidy needed per unit','zSubU')]):
    ws.cell(r+2+i,2,lab)
    for j in range(3): ws.cell(r+2+i,3+j,f"={q('New Haven')}!{L(3+j)}{ROWS[k]}").number_format=FMT['pct' if k=='zSave' else 'usd']
ch=BarChart(); ch.type='col'; sty(ch,'Total development cost per unit',14,8)
ch.add_data(Reference(ws,min_col=2,max_col=5,min_row=r+2,max_row=r+3),from_rows=True,titles_from_data=True); ch.set_categories(Reference(ws,min_col=3,max_col=5,min_row=r+1))
color(ch.series[0],'2C6A8A'); color(ch.series[1],'A9B4BA'); ch.y_axis.numFmt='$#,##0'
ws.add_chart(ch,f'G{r}')
r+=22
# E. waterfall x3
for cty in CITIES:
    blk(r,f'E. Funding waterfall, one {cty} building (caps, not awards; deck: funding stack)')
    ws.cell(r+1,3,'Base (hidden)').font=F_B; ws.cell(r+1,4,'Amount').font=F_B
    n=ROWS['wfEnd']-ROWS['wfStart']+1
    for i in range(n):
        src=ROWS['wfStart']+i
        ws.cell(r+2+i,2,f"={q(cty)}!B{src}"); ws.cell(r+2+i,3,f"={q(cty)}!D{src}").number_format=FMT['usd']; ws.cell(r+2+i,4,f"={q(cty)}!E{src}").number_format=FMT['usd']
    ch=BarChart(); ch.type='col'; ch.grouping='stacked'; ch.overlap=100; sty(ch,f'{cty}: how far the programs could go',16,9)
    ch.add_data(Reference(ws,min_col=3,max_col=4,min_row=r+1,max_row=r+1+n),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=r+2,max_row=r+1+n))
    hide(ch.series[0]); color(ch.series[1],'2C6A8A'); ch.legend=None; ch.y_axis.numFmt='$#,##0'; ch.gapWidth=60
    ws.add_chart(ch,f'F{r}')
    r+=21
# F. city profiles (6 charts)
blk(r,"F. City profiles (deck: 'Three cities, three markets'). Series = cities.")
r+=1
PROFC=[('Homeownership and local income',[('Homeownership rate',MD['Homeownership rate']),('City median income as % of AMI',MD['medPct'])],'0%'),
 ('Household income vs. AMI targets',[('City median income',MD['Median household income (city)']),('60% AMI (family of 4)',MD['60% AMI, 4-person (model: 2 x 50% limit x 0.6)']),('80% AMI (family of 4)',MD['80% AMI, 4-person (model)'])],'$#,##0'),
 ('Asking rent vs. 60% AMI limit (monthly)',[('1BR asking',MD['Asking rent, 1BR']),('2BR asking',MD['Asking rent, 2BR']),('3BR asking',MD['Asking rent, 3BR']),('60% AMI limit, 2BR',MD['60% AMI rent limit, 2BR (gross, rent + utilities)'])],'$#,##0'),
 ('Home prices',[('Median sale, all homes',MD['Median sale price, all homes (Redfin)']),('Existing 2-unit median',MD['Median sale, existing 2-unit (CAMA)'])],'$#,##0'),
 ('Need',[('Renters below 50% AMI',MD['Renters below 50% AMI']),('Poverty rate',MD['Poverty rate'])],'0%'),
 ('Residents by race and ethnicity',[('Hispanic or Latino',MD['Hispanic or Latino (any race)']),('Black',MD['Black']),('White, non-Hispanic',MD['White, non-Hispanic'])],'0%')]
for t,items,nf in PROFC:
    ws.cell(r,2,t).font=F_B
    for j,cty in enumerate(CITIES): ws.cell(r,3+j,cty).font=F_B
    for i,(lab,mr) in enumerate(items):
        ws.cell(r+1+i,2,lab)
        for j in range(3): ws.cell(r+1+i,3+j,f"='Market Data'!{L(3+j)}{mr}").number_format=nf
    ch=BarChart(); ch.type='col'; sty(ch,t,13,7)
    ch.add_data(Reference(ws,min_col=3,max_col=5,min_row=r,max_row=r+len(items)),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=r+1,max_row=r+len(items)))
    for s_,cty in zip(ch.series,CITIES): color(s_,CCOL[cty])
    ch.y_axis.numFmt=nf; ws.add_chart(ch,f'G{r}')
    r+=15
# G. income ladder
blk(r,"G. Income ladder: household income as % of AMI (deck: 'Many can rent, few can buy')")
for j,cty in enumerate(CITIES): ws.cell(r+1,3+j,cty).font=F_B
G=[('Typical city household',lambda j:f"='Market Data'!{L(3+j)}{MD['medPct']}"),('Our buyer target',lambda j:f"={q(CITIES[j])}!$C${ROWS['ownerAMI']}"),
   ('Buyer can pay the appraised value (with rent)',None),('Buy the median home (all types)',lambda j:f"='Market Data'!{L(3+j)}{MD['Income needed to buy the median home (% of AMI)']}")]
for i,(lab,fn) in enumerate(G):
    ws.cell(r+2+i,2,lab)
    for j,cty in enumerate(CITIES):
        if fn: c=ws.cell(r+2+i,3+j,fn(j)); c.number_format='0%'
        else: slot(ws,f'{L(3+j)}{r+2+i}','CROSS()',cty,'pct0')
ch=BarChart(); ch.type='col'; sty(ch,'Household income as % of AMI',14,8)
ch.add_data(Reference(ws,min_col=3,max_col=5,min_row=r+1,max_row=r+5),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=r+2,max_row=r+5))
for s_,cty in zip(ch.series,CITIES): color(s_,CCOL[cty])
ch.y_axis.numFmt='0%'; ws.add_chart(ch,f'G{r}')
ws.cell(r+6,2,'The crossover is the owner income at which the most the buyer can pay first reaches the appraised value (solved by the explorer; blank above 200% AMI).').font=F_I
r+=20
# H. buyer pays vs subsidy
blk(r,'H. One building: what the buyer pays vs. subsidy needed (deck: prefab slide)')
for j,cty in enumerate(CITIES): ws.cell(r+1,3+j,cty).font=F_B
for i,(lab,k) in enumerate([('Buyer pays','buyerPays'),('Subsidy needed','subNeed')]):
    ws.cell(r+2+i,2,lab)
    for j,cty in enumerate(CITIES): ws.cell(r+2+i,3+j,f"={q(cty)}!$C${ROWS[k]}").number_format=FMT['usd']
ch=BarChart(); ch.type='col'; ch.grouping='stacked'; ch.overlap=100; sty(ch,'Per building: buyer pays vs. subsidy',13,8)
ch.add_data(Reference(ws,min_col=2,max_col=5,min_row=r+2,max_row=r+3),from_rows=True,titles_from_data=True); ch.set_categories(Reference(ws,min_col=3,max_col=5,min_row=r+1))
color(ch.series[0],'2C6A8A'); color(ch.series[1],'B5452F'); ch.y_axis.numFmt='$#,##0'
ws.add_chart(ch,f'G{r}')
r+=18
# I. focus city sources
FOC_ROW=r
slot(ws,f'B{r}','"I. "+FOCUS+": sources vs. the "+usdK(RF.subNeed)+" subsidy needed (focus city at export; deck: prefab slide)"',None,'txt',F_B)
ws.cell(r+1,3,'Amount').font=F_B
SRCS=[('Time To Own','RF.tto'),('DOH homebuyer subsidy','RF.hb'),('HDF / C4C grant','RF.sHdf'),('FHLBank Boston AHP','RF.sAhp'),('HTCC share','RF.sHtcc'),('New Markets Tax Credits','RF.sNmtc'),('DOH developer subsidy','RF.sDoh'),('Remaining gap','RF.remain')]
for i,(lab,e) in enumerate(SRCS):
    ws.cell(r+2+i,2,lab); slot(ws,f'C{r+2+i}',e,None)
ch=BarChart(); ch.type='bar'; sty(ch,'Focus city: sources (caps, not awards)',14,8)
ch.add_data(Reference(ws,min_col=3,min_row=r+1,max_row=r+1+len(SRCS)),titles_from_data=True); ch.set_categories(Reference(ws,min_col=2,min_row=r+2,max_row=r+1+len(SRCS)))
color(ch.series[0],'2C6A8A'); ch.legend=None; ch.x_axis.scaling.orientation='maxMin'; ch.y_axis.numFmt='$#,##0'
ws.add_chart(ch,f'G{r}')
slot(ws,'B4','"Focus city at export: "+FOCUS+". Blocks D and I use the focus city or New Haven as labeled."',None,'txt',F_I)

for s_ in wb.worksheets:
    s_.sheet_view.showGridLines=False
wb.save('sub/cmc_template.xlsx')
json.dump({"slots":SLOTS,"readmeLog":README_LOG,"snapRow":SNAP_ROW,"nscen":NSCEN},open('sub/cmc_map.json','w'))
print('slots',len(SLOTS),'rows',len(ROWS))
